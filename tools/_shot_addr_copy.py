# -*- coding: utf-8 -*-
"""v1.10.0 新增功能验证 + 截屏：
① 出售记录「收货地址」（弹窗填写 / 自动带出买家地址 / 列表列 / 搜索 / 导出）；
② 出售记录页与出库记录页的复制（拖选区块 / Ctrl+C / 右键菜单项目 / 拖选即复制开关）。
"""
import os
import shutil
import sys
import tempfile
import tkinter as tk

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))

import wm.app as wapp  # noqa: E402
from wm.sales import SaleDialog  # noqa: E402
from wm.store import Store, blank_record  # noqa: E402
import screenshot  # noqa: E402

TMP = tempfile.mkdtemp(prefix="wm_addr_")
wapp.app_dir = lambda: TMP
OUT = os.path.join(ROOT, ".tmpdbg")
os.makedirs(OUT, exist_ok=True)

FAILS = []


def check(name, ok, extra=""):
    print(("  [PASS] " if ok else "  [FAIL] ") + name + ((" -> " + str(extra)) if extra and not ok else ""))
    if not ok:
        FAILS.append(name)


boot = Store(os.path.join(TMP, "data", "warehouse.db"))
boot.close()

root = tk.Tk()
root.withdraw()
app = wapp.App(root)

data = [
    ("WL-00001", "罗技", "外设", "MX MASTER 3", "99新", 599.0, 4, "6排A-3", 1),
    ("WL-00002", "Apple", "手机", "A2634", "95新", 2899.0, 3, "7排B-2", 1),
    ("WL-00003", "戴尔", "显示器", "U2720Q", "9成新", 1580.0, 2, "8排C-1", 1),
]
app.store.insert_many([blank_record() | {"code": c, "brand": b, "category": cat,
                                         "model": m, "condition": cond, "price": p,
                                         "qty": q, "shelf": sh, "listed": l}
                       for c, b, cat, m, cond, p, q, sh, l in data])
app.load_data()
recs = {r["code"]: r for r in app.store.all_records()}
app.store.add_sale(recs["WL-00001"]["id"], 1, 650.0, buyer="张三", sale_date="2026-10-02",
                   address="广东省深圳市南山区科技园南路 1 号 A 座 1201", remark="微信转账")
app.store.add_sale(recs["WL-00001"]["id"], 1, 620.0, buyer="张三", sale_date="2026-10-05",
                   address="广东省深圳市南山区科技园南路 1 号 A 座 1201")
app.store.add_sale(recs["WL-00002"]["id"], 1, 3100.0, buyer="李四", sale_date="2026-10-06",
                   address="北京市朝阳区建国路 88 号 SOHO 现代城 B 座")
app.store.add_sale(recs["WL-00003"]["id"], 1, 0, sale_date="2026-10-07", out_type="报废")
app.load_data()

print("== 1. 数据层：地址写入 / 搜索 / 买家地址映射 ==")
rows = app.store.list_sales(scope="sale")
check("销售记录存了地址", all(r["address"] for r in rows),
      str([(r["buyer"], r["address"]) for r in rows]))
check("按地址关键词搜索", len(app.store.list_sales(kw="南山")) == 2,
      str(len(app.store.list_sales(kw="南山"))))
check("按城市搜索", len(app.store.list_sales(kw="北京")) == 1)
check("非销售出库不记地址",
      all(not r["address"] for r in app.store.list_sales(scope="out")))
check("买家最近地址映射",
      app.store.buyer_address_map().get("李四") == "北京市朝阳区建国路 88 号 SOHO 现代城 B 座",
      str(app.store.buyer_address_map()))
check("历史地址池（最近在前，去重）",
      len(app.store.sale_addresses()) == 2, str(app.store.sale_addresses()))

print("== 2. 登记出售弹窗：收货地址行 + 自动带出 ==")
d = SaleDialog(root, recs["WL-00002"], app.fam,
               buyers=app.store.buyer_names(),
               addr_map=app.store.buyer_address_map(), mode="sale")
d.update()
check("弹窗有「收货地址」行", "收货地址" in d.hints, str(list(d.hints)))
check("地址下拉含历史地址", len(getattr(d, "addr_pool", [])) == 2, str(d.addr_pool))
d.var_buyer.set("李四")
d.var_addr.set("")
d._fill_addr()
check("选中买家自动带出最近地址",
      d.var_addr.get() == "北京市朝阳区建国路 88 号 SOHO 现代城 B 座", d.var_addr.get())
d.var_qty.set("1")
d.var_price.set("3100")
d.attributes("-topmost", True)
d.lift()
d.update()
screenshot.capture(os.path.join(OUT, "addr_dialog.png"))
print("shot -> addr_dialog.png")
d._cancel()
check("取消后无结果", d.result is None)

d2 = SaleDialog(root, recs["WL-00003"], app.fam, buyers=app.store.buyer_names(),
                addr_map=app.store.buyer_address_map(), mode="sale")
d2.update()
d2.var_qty.set("1")
d2.var_price.set("1500")
d2.var_buyer.set("王五")
d2.var_addr.set("上海市浦东新区世纪大道 100 号")
d2._ok()
check("弹窗返回结果含地址",
      d2.result and d2.result["address"] == "上海市浦东新区世纪大道 100 号",
      str(d2.result))
app.store.add_sale(recs["WL-00003"]["id"], d2.result["qty"], d2.result["price"],
                   d2.result["buyer"], d2.result["sale_date"], d2.result["remark"],
                   d2.result["out_type"], d2.result["address"])
app.load_data()

d3 = SaleDialog(root, recs["WL-00003"], app.fam, mode="out")
d3.update()
check("登记出库弹窗没有地址行", "收货地址" not in d3.hints, str(list(d3.hints)))
d3._cancel()

print("== 3. 出售记录页：地址列 ==")
root.deiconify()
root.geometry("1500x880+60+30")
app.nb.select(1)
root.update()
tb = app.tab_sales
tb.refresh_tab()
root.update()
keys = tb.col_keys()
check("列表含地址列且位置在买家之后",
      keys.index("address") == keys.index("buyer") + 1, str(keys))
kid_addr = next(k for k in tb.tree.get_children() if tb.row_map()[k].get("address"))
vals = tb.tree.item(kid_addr, "values")
check("列表显示地址文本", any("广东省" in v or "北京市" in v or "上海市" in v
                              for v in vals), str(vals))

print("== 4. 复制：拖选区块 / Ctrl+C / 右键项目 ==")
sx = keys.index("code")
ex = keys.index("address")
tb._drag_block = {"rows": [kid_addr], "c0": sx, "c1": ex}
tb.copy_block()
blk = root.clipboard_get()
check("拖选区块复制为 TSV", len(blk.split("\t")) == ex - sx + 1, repr(blk[:80]))
check("区块内容含地址", any(x in blk for x in ("广东省", "北京市", "上海市")), repr(blk[:80]))
tb._drag_block = None
tb.tree.selection_set(kid_addr)
tb.tree.focus(kid_addr)
tb._cell_key = "address"
root.update()
tb.copy_cell()
check("复制单元格＝该行地址",
      root.clipboard_get() == tb.row_map()[kid_addr]["address"], root.clipboard_get())
tb.copy_row()
check("复制整行（字段: 值）", "收货地址：" in root.clipboard_get(), root.clipboard_get()[:60])
tb.copy_shortcut()
check("Ctrl+C 复制选中行", "\t" in root.clipboard_get())
tb.copy_with_header()
check("复制表头 + 全部行", root.clipboard_get().split("\n")[0].count("\t") == len(keys) - 1)

# 拖选区块视觉状态（截图用）
tb._drag_block = {"rows": list(tb.tree.get_children())[:2],
                  "c0": keys.index("code"), "c1": keys.index("buyer")}
tb._status("已选中区块：2 行 × 4 列（编号 ~ 买家/客户），按 Ctrl+C 复制")
root.update()
screenshot.capture_window("v1.10.0", os.path.join(OUT, "sales_addr.png"))
print("shot -> sales_addr.png（出售记录页）")
tb._drag_block = None

print("== 5. 出库记录页复制 ==")
app.nb.select(2)
root.update()
ob = app.tab_out
ob.refresh_tab()
root.update()
okeys = ob.col_keys()
okids = ob.tree.get_children()
ob._drag_block = {"rows": list(okids[:2]), "c0": 0, "c1": okeys.index("qty")}
ob.copy_block()
check("出库页拖选区块复制", root.clipboard_get().count("\n") == 1,
      repr(root.clipboard_get()[:80]))
ob._drag_block = None
ob.tree.selection_set(okids[0])
ob.tree.focus(okids[0])
root.update()
ob.copy_shortcut()
check("出库页 Ctrl+C", "\t" in root.clipboard_get())
screenshot.capture_window("v1.10.0", os.path.join(OUT, "out_copy.png"))
print("shot -> out_copy.png（出库记录页）")

print("== 6. 导出 XLSX 含地址列 ==")
import openpyxl  # noqa: E402
from wm.xlsxio import export_sales_xlsx  # noqa: E402
xp = os.path.join(TMP, "sale_addr.xlsx")
export_sales_xlsx(xp, app.store.list_sales(scope="sale"), "地址测试")
wb = openpyxl.load_workbook(xp)
ws = wb["出售记录"]
head = [c.value for c in ws[1]]
check("导出表头含「收货地址」", "收货地址" in head, str(head))
ci = head.index("收货地址") + 1
check("导出地址值正确",
      ws.cell(row=2, column=ci).value in [r["address"] for r in app.store.list_sales(scope="sale")],
      str(ws.cell(row=2, column=ci).value))
wb.close()

print()
if FAILS:
    print("FAILURES: %d -> %s" % (len(FAILS), FAILS))
else:
    print("ALL CHECKS PASSED")

app._stopped = True
try:
    root.destroy()
except Exception:
    pass
shutil.rmtree(TMP, ignore_errors=True)
