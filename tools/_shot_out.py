# -*- coding: utf-8 -*-
"""临时脚本：出售/出库分页验证 + 截图（v1.9.3 出库与销售分开）。"""
import os
import shutil
import sys
import tempfile
import tkinter as tk

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))

import wm.app as wapp  # noqa: E402
from wm.sales import SaleDialog  # noqa: E402
from wm.store import Store, blank_record  # noqa: E402
import screenshot  # noqa: E402

TMP = tempfile.mkdtemp(prefix="wm_out_")
wapp.app_dir = lambda: TMP
OUT = os.path.join(ROOT, ".tmpdbg")
os.makedirs(OUT, exist_ok=True)

st = Store(os.path.join(TMP, "data", "warehouse.db"))
st.close()

root = tk.Tk()
root.withdraw()
app = wapp.App(root)

DATA = [
    ("WL-00001", "罗技", "外设", "MX MASTER 3", 599.0, 8),
    ("WL-00002", "Apple", "手机", "A2634", 2899.0, 4),
    ("WL-00003", "戴尔", "显示器", "U2720Q", 1580.0, 3),
    ("WL-00004", "华为", "网络", "S5720", 2100.0, 5),
]
app.store.insert_many([blank_record() | {"code": c, "brand": b, "category": cat,
                                         "model": m, "price": p, "qty": q, "listed": 1}
                       for c, b, cat, m, p, q in DATA])
recs = {r["code"]: r for r in app.store.all_records()}
# 销售 2 笔
app.store.add_sale(recs["WL-00001"]["id"], 1, 650.0, buyer="张三", sale_date="2026-10-02")
app.store.add_sale(recs["WL-00002"]["id"], 2, 2900.0, buyer="李四", sale_date="2026-10-05")
# 非销售出库 3 笔
app.store.add_sale(recs["WL-00003"]["id"], 1, 0, sale_date="2026-10-05",
                   out_type="报废", remark="屏幕碎裂")
app.store.add_sale(recs["WL-00003"]["id"], 1, 0, sale_date="2026-10-06",
                   out_type="退货", remark="退回供货商")
app.store.add_sale(recs["WL-00004"]["id"], 1, 0, sale_date="2026-10-06",
                   out_type="借出", remark="借给门店试用")
app.load_data()
root.title("仓管家-出库验证窗口")
root.attributes("-topmost", True)
root.lift()
root.deiconify()
root.geometry("1440x900+60+20")
root.update()

FAIL = []


def check(name, cond, extra=""):
    print("  [%s] %s %s" % ("PASS" if cond else "FAIL", name, extra if not cond else ""))
    if not cond:
        FAIL.append(name)


print("== 1. 页面分流 ==")
check("共 4 个标签页", len(app.nb.tabs()) == 4, str(app.nb.tabs()))
check("第 2 页=出售记录 第 3 页=出库记录 第 4 页=数据分析",
      "出售记录" in app.nb.tab(1, "text") and "出库记录" in app.nb.tab(2, "text")
      and "数据分析" in app.nb.tab(3, "text"),
      " | ".join(app.nb.tab(i, "text") for i in range(4)))
app.nb.select(1)
app.tab_sales.refresh_tab()
root.update()
check("出售页只有销售记录（2 笔）", len(app.tab_sales.rows) == 2,
      str(len(app.tab_sales.rows)))
check("出售页销售额 = 6450", abs(app.store.sales_totals(app.tab_sales.rows)["amount"]
                                 - 6450.0) < 1e-6)
check("出售页不含非销售出库行",
      all((r.get("out_type") or "销售") == "销售" for r in app.tab_sales.rows),
      str([r.get("out_type") for r in app.tab_sales.rows]))
screenshot.capture_window("出库验证", os.path.join(OUT, "out_tab_sale.png"))
print("shot -> out_tab_sale.png（出售记录页）")

app.nb.select(2)
app.tab_out.refresh_tab()
root.update()
check("出库页总览全部出库（2 销售 + 3 非销售 = 5 笔）",
      len(app.tab_out.rows) == 5, str(len(app.tab_out.rows)))
check("出库页含销售行与报废行",
      {(r.get("out_type") or "销售") for r in app.tab_out.rows}
      == {"销售", "报废", "退货", "借出"},
      str({(r.get("out_type") or "销售") for r in app.tab_out.rows}))
check("出库页无售价列", "price" not in app.tab_out.tree["columns"],
      str(app.tab_out.tree["columns"]))
check("出库页有类型/成本列",
      "out_type" in app.tab_out.tree["columns"]
      and "cost_total" in app.tab_out.tree["columns"])
t = app.store.out_totals(app.tab_out.rows)
check("出库页 KPI 笔数/件数（5 笔 / 6 件）",
      app.tab_out.kpi["count"].cget("text") == "5"
      and app.tab_out.kpi["qty"].cget("text") == "6",
      "%s/%s" % (app.tab_out.kpi["count"].cget("text"),
                 app.tab_out.kpi["qty"].cget("text")))
check("出库页 KPI 销售/非销售拆分（2 / 3 笔）",
      app.tab_out.kpi["split"].cget("text") == "2 / 3 笔",
      app.tab_out.kpi["split"].cget("text"))
check("出库页类型筛选存在且含销售",
      app.tab_out.cb_type is not None
      and "销售" in app.tab_out.cb_type.cget("values"))
check("出库页无销售目标", not hasattr(app.tab_out, "bar_target"))
check("出库页提示文案含销售与非销售",
      "销售 2 笔" in app.tab_out.lbl_hint.cget("text")
      and "非销售 3 笔" in app.tab_out.lbl_hint.cget("text"),
      app.tab_out.lbl_hint.cget("text"))
screenshot.capture_window("出库验证", os.path.join(OUT, "out_tab.png"))
print("shot -> out_tab.png（出库记录页）")

print("== 2. 两个登记弹窗 ==")
d_sale = SaleDialog(root, recs["WL-00001"], app.fam, buyers=["张三"], mode="sale")
d_sale.update()
check("出售弹窗：无类型行/有售价行", "出库类型" not in d_sale.hints
      and "出售单价" in d_sale.hints)
d_sale.var_qty.set("1")
d_sale.var_price.set("650")
d_sale._ok()
check("出售弹窗结果 out_type=销售", d_sale.result["out_type"] == "销售", str(d_sale.result))

d_out = SaleDialog(root, recs["WL-00001"], app.fam, mode="out")
d_out.update()
check("出库弹窗：有类型行/无售价行", "出库类型" in d_out.hints
      and "出售单价" not in d_out.hints)
check("出库弹窗：默认类型非销售", d_out.var_type.get() != "销售",
      d_out.var_type.get())
d_out.var_qty.set("2")
d_out._ok()
check("出库弹窗结果：价格为 0、无买家", d_out.result["price"] == 0.0
      and d_out.result["buyer"] == "", str(d_out.result))

print("== 3. 出库页类型筛选 ==")
app.tab_out.var_type.set("报废")
app.tab_out.refresh_rows()
check("按类型筛选-报废 1 笔", len(app.tab_out.rows) == 1, str(len(app.tab_out.rows)))
check("筛选后 KPI 跟着变（1 笔 / 1 件）",
      app.tab_out.kpi["count"].cget("text") == "1"
      and app.tab_out.kpi["qty"].cget("text") == "1",
      "%s/%s" % (app.tab_out.kpi["count"].cget("text"),
                 app.tab_out.kpi["qty"].cget("text")))
app.tab_out.var_type.set("销售")
app.tab_out.refresh_rows()
check("按类型筛选-销售 2 笔", len(app.tab_out.rows) == 2, str(len(app.tab_out.rows)))
app.tab_out.var_type.set("全部")
app.tab_out.refresh_rows()
root.update()

print("== 4. 出库导出 XLSX（总览）==")
import openpyxl  # noqa: E402
from wm.xlsxio import export_out_xlsx  # noqa: E402
p_out = os.path.join(OUT, "out_export.xlsx")
export_out_xlsx(p_out, app.tab_out.rows, "测试")
wbo = openpyxl.load_workbook(p_out)
check("出库导出工作表", "出库记录" in wbo.sheetnames and "出库汇总" in wbo.sheetnames,
      str(wbo.sheetnames))
hdr = [c.value for c in wbo["出库记录"][1]]
check("出库导出表头无售价/销售额", "售价" not in hdr and "销售额" not in hdr
      and "利润" not in hdr, str(hdr))
check("出库导出表头含类型/成本",
      hdr[:4] == ["出库日期", "类型", "编号", "品牌"], str(hdr[:4]))
check("出库导出行数 = 全部 5 笔",
      wbo["出库记录"].max_row == 6, str(wbo["出库记录"].max_row))
check("出库导出含销售行",
      any(c.value == "销售" for c in wbo["出库记录"]["B"]),
      str([c.value for c in wbo["出库记录"]["B"]]))
wbo.close()

root.destroy()
shutil.rmtree(TMP, ignore_errors=True)
print("\n%s" % ("ALL PASSED" if not FAIL else "FAILED: %s" % FAIL))
