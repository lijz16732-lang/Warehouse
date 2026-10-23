# -*- coding: utf-8 -*-
"""XLSX 读写：智能识别表头、三种导入模式、样式化导出。"""

import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import (FIELD_KEYS, FIELD_LABEL, FIELD_TYPE, FIELD_WIDTH,
                        HEADER_ALIASES, LISTED_FALSE, LISTED_TRUE, LISTED_TEXT,
                        MULTI_SEP, OUT_TYPE_SALE)


# ---------------------------------------------------------------- 归一化
_PUNCT = re.compile(r"[\s\u3000()（）\[\]【】:：,，.。/\\\-_*#·'\"`~!！?？+＋=＝|]+")
_PAREN = re.compile(r"[（(]([^()（）]*)[)）]|【([^【】]*)】|\[([^\[\]]*)\]")


def norm_header(text):
    """表头归一化：全角转半角、去标点空格、转小写。"""
    if text is None:
        return ""
    s = unicodedata.normalize("NFKC", str(text)).strip().lower()
    s = _PUNCT.sub("", s)
    # 去掉常见单位后缀
    s = re.sub(r"(元|rmb|cny|个|件|台|只|支|套|人民币)$", "", s)
    return s


def header_candidates(text):
    """由原始表头生成候选名（按优先级）：
    1. 整体归一化        「G (货架)」→ g货架
    2. 括号内文本        「G (货架)」→ 货架
    3. 括号外文本        「货架(旧)」→ 货架
    """
    s = "" if text is None else str(text)
    out = [norm_header(s)]
    m = _PAREN.search(s)
    if m:
        inner = next(g for g in m.groups() if g is not None)
        out.append(norm_header(inner))
        out.append(norm_header(s[:m.start()] + s[m.end():]))
    return [c for c in out if c]


_ALIAS_SET = {}
for _k, _alist in HEADER_ALIASES.items():
    for _a in _alist:
        _ALIAS_SET.setdefault(norm_header(_a), set()).add(_k)


def _match_alias(cands, keys=None):
    """候选名精确匹配别名，返回字段 key 或 None。"""
    for c in cands:
        hits = _ALIAS_SET.get(c)
        if hits:
            for k in FIELD_KEYS:
                if k in hits and (keys is None or k in keys):
                    return k
    return None


def _seq_suffix(c, a):
    """候选名 c 是否为「别名 a + 纯数字序号」（如 位置1 / 货架2 / shelf3）。
    这类列表示同一字段的第 N 份数据（多货架等），允许多列映射到同字段。"""
    if c.startswith(a):
        rest = c[len(a):]
    elif c.endswith(a):
        rest = c[:-len(a)]
    else:
        return False
    return bool(re.fullmatch(r"\d{1,3}", rest))


def map_headers(headers):
    """把表头列表映射成 {列下标: 字段key}。
    精确匹配（含括号内/外候选）优先，其次包含匹配。
    文本字段允许「位置1/位置2」这类序号列多列映射到同字段（导入时合并）。"""
    mapping, used = {}, set()
    cands = [header_candidates(h) for h in headers]

    # 第一轮：精确匹配
    for i, cs in enumerate(cands):
        k = _match_alias(cs)
        if k and k not in used:
            mapping[i] = k
            used.add(k)

    # 第二轮：包含匹配（别名牌头里包含字段名，或字段名包含别名牌头）；
    # 文本字段允许「货架2 / 位置3」这类序号列重复映射（多货架多列合并）
    for i, cs in enumerate(cands):
        if i in mapping:
            continue
        for c in cs:
            if len(c) < 2:
                break
            for k in FIELD_KEYS:
                if k in used and not (FIELD_TYPE.get(k) == "text"
                                      and any(_seq_suffix(c, norm_header(a))
                                              for a in HEADER_ALIASES[k])):
                    continue
                for alias in HEADER_ALIASES[k]:
                    a = norm_header(alias)
                    if a and len(a) >= 2 and (a in c or c in a):
                        mapping[i] = k
                        used.add(k)
                        break
                if i in mapping:
                    break
            if i in mapping:
                break
    # 第三轮：近似匹配（同长度、首字相同的短表头，容忍「单格/单价」这类错别字或简写）
    for i, cs in enumerate(cands):
        if i in mapping:
            continue
        for c in cs:
            if not (2 <= len(c) <= 3):
                continue
            for k in FIELD_KEYS:
                if k in used:
                    continue
                hit = False
                for alias in HEADER_ALIASES[k]:
                    a = norm_header(alias)
                    if len(a) == len(c) and a[:1] == c[:1]:
                        hit = True
                        break
                if hit:
                    mapping[i] = k
                    used.add(k)
                    break
            if i in mapping:
                break
    return mapping


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def to_bool(v):
    if isinstance(v, bool):
        return 1 if v else 0
    if isinstance(v, (int, float)):
        return 1 if v else 0
    s = _cell_text(v)
    return 1 if s in LISTED_TRUE else 0


def to_float(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return float(s)
    except ValueError:
        return 0.0


def to_int(v):
    """转整数；「1+1」这类累加写法按求和处理（=2）。"""
    s = _cell_text(v)
    if "+" in s:
        parts = [to_float(p) for p in s.split("+") if p.strip()]
        if parts:
            return int(round(sum(parts)))
    try:
        return int(round(to_float(v)))
    except (TypeError, ValueError):
        return 0


# ---------------------------------------------------------------- 读取
def list_sheets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def detect_header_row(ws, max_scan=15):
    """在前若干行里挑出最像表头的一行（别名命中数优先，同行数取更靠前的行）。"""
    best_row, best_score = 1, 0
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        cells = [_cell_text(c) for c in row]
        non_empty = [c for c in cells if c]
        if not non_empty:
            continue
        hits = sum(1 for c in non_empty if _match_alias(header_candidates(c)))
        score = hits * 3
        score += len(non_empty) if len(non_empty) >= 3 else 0
        if score > best_score:
            best_score, best_row = score, idx
    return best_row


def read_xlsx(path, sheet=None, header_row=None):
    """读取 xlsx，返回 (headers, rows, meta)。rows 为 list[list]。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    try:
        if header_row is None:
            header_row = detect_header_row(ws)
        raw = list(ws.iter_rows(min_row=1, max_row=min(header_row + 1, ws.max_row or 1),
                                values_only=True))
        headers = [_cell_text(c) for c in (raw[header_row - 1] if len(raw) >= header_row else [])]
        rows = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if row is None:
                continue
            if all(c is None or _cell_text(c) == "" for c in row):
                continue
            rows.append(list(row))
    finally:
        wb.close()
    meta = {"sheet": ws.title, "sheetnames": list_sheets(path), "header_row": header_row}
    return headers, rows, meta


def rows_to_records(headers, rows, mapping=None):
    """按表头映射把原始行转成记录字典，返回 (records, mapping, 未识别表头)。
    可传入手动指定的 mapping 覆盖自动识别。
    同一字段映射多列时（多货架等），非空值按列顺序用「 / 」合并。"""
    if mapping is None:
        mapping = map_headers(headers)
    unmapped = [headers[i] for i in range(len(headers))
                if i not in mapping and _cell_text(headers[i])]
    records = []
    ordered = sorted(mapping.items())   # 按列顺序合并，保证「位置1 / 位置2」次序稳定
    for row in rows:
        rec = {}
        for i, field in ordered:
            if i >= len(row):
                continue
            v = row[i]
            if field in rec:
                prev, cur = _cell_text(rec[field]), _cell_text(v)
                if cur and cur != prev:
                    rec[field] = (prev + MULTI_SEP + cur) if prev else cur
            else:
                rec[field] = v
        if any(_cell_text(v) for v in rec.values()):
            records.append(rec)
    return records, mapping, unmapped


def normalize_record(raw):
    """把原始单元格值转成入库类型。"""
    out = {}
    for k in FIELD_KEYS:
        t = FIELD_TYPE[k]
        v = raw.get(k)
        if t == "money":
            out[k] = to_float(v)
        elif t == "int":
            out[k] = to_int(v)
        elif t == "bool":
            out[k] = to_bool(v) if k in raw else 0
        else:
            out[k] = _cell_text(v)
    return out


# ---------------------------------------------------------------- 写出
_HDR_FILL = PatternFill("solid", fgColor="1D4ED8")
_HDR_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
_BODY_FONT = Font(name="微软雅黑", size=10)
_ALT_FILL = PatternFill("solid", fgColor="F4F7FC")
_THIN = Side(style="thin", color="D6DEEA")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_sheet(ws, headers, n_rows, widths=None):
    for c, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 14)
    ws.row_dimensions[1].height = 26
    for r in range(2, n_rows + 2):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = _BODY_FONT
            cell.border = _BORDER
            if r % 2 == 0:
                cell.fill = _ALT_FILL
    ws.freeze_panes = "A2"


def export_xlsx(path, records, columns=None, sheet_name="库存明细"):
    """导出记录列表为带样式的 xlsx。"""
    columns = columns or FIELD_KEYS
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    headers = [FIELD_LABEL[c] for c in columns]
    widths = {i + 1: FIELD_WIDTH[c] / 8.2 + 2 for i, c in enumerate(columns)}

    ws.append(headers)
    for rec in records:
        row = []
        for c in columns:
            v = rec.get(c)
            t = FIELD_TYPE[c]
            if t == "bool":
                row.append(LISTED_TEXT.get(int(v or 0), "未上架"))
            elif t == "money":
                row.append(round(float(v or 0), 2))
            elif t == "int":
                row.append(int(v or 0))
            else:
                row.append("" if v is None else str(v))
        ws.append(row)

    _style_sheet(ws, headers, len(records), widths)

    for i, c in enumerate(columns, start=1):
        letter = get_column_letter(i)
        if FIELD_TYPE[c] == "money":
            for r in range(2, len(records) + 2):
                ws["%s%d" % (letter, r)].number_format = "#,##0.00"
                ws["%s%d" % (letter, r)].alignment = Alignment(horizontal="right")
        elif FIELD_TYPE[c] == "int":
            for r in range(2, len(records) + 2):
                ws["%s%d" % (letter, r)].number_format = "#,##0"
                ws["%s%d" % (letter, r)].alignment = Alignment(horizontal="right")
        elif FIELD_TYPE[c] == "bool":
            for r in range(2, len(records) + 2):
                ws["%s%d" % (letter, r)].alignment = Alignment(horizontal="center")
        elif FIELD_ANCHOR_OF(c) == "center":
            for r in range(2, len(records) + 2):
                ws["%s%d" % (letter, r)].alignment = Alignment(horizontal="center")

    if records:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(columns)), len(records) + 1)
    wb.save(path)
    return path


def FIELD_ANCHOR_OF(key):
    from .constants import FIELD_ANCHOR
    return FIELD_ANCHOR.get(key, "w")


def export_stats_xlsx(path, records, stats, scope_text=""):
    """导出「明细 + 多维度统计」多工作表报表。"""
    export_xlsx(path, records, FIELD_KEYS, "库存明细")

    wb = load_workbook(path)
    ws = wb.create_sheet("统计概览")

    def put(row, a, b=None, bold=False, fill=None):
        ws.cell(row=row, column=1, value=a)
        if b is not None:
            ws.cell(row=row, column=2, value=b)
        for c in (1, 2):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(name="微软雅黑", size=10.5, bold=bold)
            cell.border = _BORDER
            if fill:
                cell.fill = fill

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20
    for col in ("C", "D", "E"):
        ws.column_dimensions[col].width = 16

    r = 1
    put(r, "指标", "数值", bold=True, fill=_HDR_FILL)
    for c in (1, 2):
        ws.cell(row=r, column=c).font = _HDR_FONT
    r += 1
    rows = [
        ("统计范围", scope_text or "全部数据"),
        ("记录数（SKU）", stats["total"]),
        ("库存总数量", stats["total_qty"]),
        ("库存总价值(元)", round(stats["total_value"], 2)),
        ("已上架", stats["listed"]),
        ("未上架", stats["unlisted"]),
        ("上架率", "%.1f%%" % stats["listed_rate"]),
        ("平均价格(元)", round(stats["avg_price"], 2)),
        ("最高价格(元)", round(stats["max_price"], 2)),
        ("最低价格(元)", round(stats["min_price"], 2)),
        ("品牌数", stats["brand_count"]),
        ("品类数", stats["category_count"]),
        ("货架数", stats["shelf_count"]),
    ]
    for k, v in rows:
        put(r, k, v)
        r += 1

    def dim_sheet(name, title, data):
        s = wb.create_sheet(name)
        heads = [title, "记录数", "库存数量", "库存价值(元)", "占比"]
        for i, h in enumerate(heads, start=1):
            c = s.cell(row=1, column=i, value=h)
            c.fill, c.font = _HDR_FILL, _HDR_FONT
            c.alignment = Alignment(horizontal="center")
            c.border = _BORDER
        s.column_dimensions["A"].width = 20
        for col, w in (("B", 12), ("C", 12), ("D", 16), ("E", 10)):
            s.column_dimensions[col].width = w
        tot = stats["total"] or 1
        for i, (k, _, slot) in enumerate(data, start=2):
            s.cell(row=i, column=1, value=k)
            s.cell(row=i, column=2, value=slot["count"])
            s.cell(row=i, column=3, value=slot["qty"])
            s.cell(row=i, column=4, value=round(slot["value"], 2))
            s.cell(row=i, column=5, value="%.1f%%" % (slot["count"] / tot * 100))
            for c in range(1, 6):
                cell = s.cell(row=i, column=c)
                cell.font = _BODY_FONT
                cell.border = _BORDER
                if c == 4:
                    cell.number_format = "#,##0.00"
        s.freeze_panes = "A2"

    dim_sheet("品牌统计", "品牌", stats["brands"])
    dim_sheet("品类统计", "品类", stats["categories"])
    dim_sheet("成色统计", "成色", stats["conditions"])
    dim_sheet("货架统计", "货架", stats["shelves"])

    s = wb.create_sheet("价格区间")
    s.cell(row=1, column=1, value="价格区间(元)")
    s.cell(row=1, column=2, value="记录数")
    for c in (1, 2):
        cell = s.cell(row=1, column=c)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER
    s.column_dimensions["A"].width = 18
    s.column_dimensions["B"].width = 12
    for i, (label, cnt) in enumerate(stats["price_dist"], start=2):
        s.cell(row=i, column=1, value=label).font = _BODY_FONT
        s.cell(row=i, column=2, value=cnt).font = _BODY_FONT
    s.freeze_panes = "A2"

    wb.save(path)
    return path


SALE_HEADERS = (
    ("sale_date", "出售日期", 13), ("code", "编号", 13),
    ("brand", "品牌", 16), ("category", "品类", 12), ("model", "型号", 20),
    ("shelf", "货架", 14), ("qty", "数量", 9), ("price", "售价(元)", 11),
    ("total", "总额(元)", 12), ("cost", "成本单价(元)", 13), ("profit", "利润(元)", 11),
    ("buyer", "买家/客户", 15), ("address", "收货地址", 30), ("remark", "备注", 22),
)


def export_sales_xlsx(path, rows, scope_text=""):
    """导出「出售记录」：明细表 + 汇总表（分品牌与合计）。

    只含销售口径字段（售价/销售额/成本/利润/买家）；非销售出库请看
    export_out_xlsx（出库记录总览）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "出售记录"
    headers = [h[1] for h in SALE_HEADERS]
    widths = {i + 1: h[2] for i, h in enumerate(SALE_HEADERS)}
    ws.append(headers)
    for r in rows:
        qty = int(r.get("qty") or 0)
        total = float(r.get("total") or 0)
        cost = float(r.get("cost") or 0)
        ws.append([r.get("sale_date") or "", r.get("code") or "",
                   r.get("brand") or "", r.get("category") or "",
                   r.get("model") or "", r.get("shelf") or "",
                   qty, round(float(r.get("price") or 0), 2), round(total, 2),
                   round(cost, 2), round(total - cost * qty, 2),
                   r.get("buyer") or "", r.get("address") or "",
                   r.get("remark") or ""])
    _style_sheet(ws, headers, len(rows), widths)
    for i in range(2, len(rows) + 2):   # 数量列右对齐
        ws["G%d" % i].alignment = Alignment(horizontal="right")
    for idx in (8, 9, 10, 11):
        letter = get_column_letter(idx)
        for i in range(2, len(rows) + 2):
            ws["%s%d" % (letter, i)].number_format = "#,##0.00"
            ws["%s%d" % (letter, i)].alignment = Alignment(horizontal="right")
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)

    # ---- 汇总表：分品牌 + 合计 ----
    s = wb.create_sheet("出售汇总")
    s.cell(row=1, column=1, value="指标").font = _HDR_FONT
    s.cell(row=1, column=1).fill = _HDR_FILL
    s.cell(row=1, column=2, value="数值").font = _HDR_FONT
    s.cell(row=1, column=2).fill = _HDR_FILL
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 20
    s.cell(row=2, column=1, value="统计范围").font = _BODY_FONT
    s.cell(row=2, column=2, value=scope_text or "全部").font = _BODY_FONT

    sale_rows = list(rows)
    qty_all = sum(int(r.get("qty") or 0) for r in sale_rows)
    amt_all = sum(float(r.get("total") or 0) for r in sale_rows)
    cost_all = sum(float(r.get("cost") or 0) * int(r.get("qty") or 0) for r in sale_rows)
    r0 = 3
    for label, val in (("出售笔数", len(sale_rows)), ("出售件数", qty_all),
                       ("销售额(元)", round(amt_all, 2)), ("成本(元)", round(cost_all, 2)),
                       ("利润(元)", round(amt_all - cost_all, 2))):
        s.cell(row=r0, column=1, value=label).font = _BODY_FONT
        c = s.cell(row=r0, column=2, value=val)
        c.font = _BODY_FONT
        if isinstance(val, float):
            c.number_format = "#,##0.00"
        r0 += 1

    brand = {}
    for r in sale_rows:
        k = (r.get("brand") or "").strip() or "（无品牌）"
        slot = brand.setdefault(k, {"count": 0, "qty": 0, "amount": 0.0, "cost": 0.0})
        slot["count"] += 1
        slot["qty"] += int(r.get("qty") or 0)
        slot["amount"] += float(r.get("total") or 0)
        slot["cost"] += float(r.get("cost") or 0) * int(r.get("qty") or 0)
    r0 += 1
    s.cell(row=r0, column=1, value="分品牌统计").font = Font(
        name="微软雅黑", size=10.5, bold=True)
    r0 += 1
    for i, h in enumerate(("品牌", "笔数", "件数", "销售额(元)", "利润(元)"), start=1):
        c = s.cell(row=r0, column=i, value=h)
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(horizontal="center")
        c.border = _BORDER
    for col, w in (("A", 26), ("B", 12), ("C", 12), ("D", 16), ("E", 16)):
        s.column_dimensions[col].width = w
    for k in sorted(brand, key=lambda x: -brand[x]["amount"]):
        v = brand[k]
        r0 += 1
        s.cell(row=r0, column=1, value=k).font = _BODY_FONT
        s.cell(row=r0, column=2, value=v["count"]).font = _BODY_FONT
        s.cell(row=r0, column=3, value=v["qty"]).font = _BODY_FONT
        for col, val in ((4, round(v["amount"], 2)), (5, round(v["amount"] - v["cost"], 2))):
            c = s.cell(row=r0, column=col, value=val)
            c.font = _BODY_FONT
            c.number_format = "#,##0.00"
    wb.save(path)
    return path


OUT_HEADERS = (
    ("out_date", "出库日期", 13), ("out_type", "类型", 10), ("code", "编号", 13),
    ("brand", "品牌", 16), ("category", "品类", 12), ("model", "型号", 20),
    ("shelf", "货架", 14), ("qty", "数量", 9),
    ("unit_cost", "成本单价(元)", 13), ("cost_total", "成本合计(元)", 13),
    ("remark", "原因/去向", 22),
)


def export_out_xlsx(path, rows, scope_text=""):
    """导出「出库记录」总览：明细表 + 分类型汇总。

    出库 = 全部减库存的流水（销售出库 + 非销售出库）。本表统一用
    「库存减少价值（成本单价 × 数量）」口径，不出现售价/销售额/利润，
    以免与出售记录混淆；销售口径请看 export_sales_xlsx（出售记录）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "出库记录"
    headers = [h[1] for h in OUT_HEADERS]
    widths = {i + 1: h[2] for i, h in enumerate(OUT_HEADERS)}
    ws.append(headers)
    for r in rows:
        qty = int(r.get("qty") or 0)
        cost = float(r.get("cost") or 0)
        ws.append([r.get("sale_date") or "", r.get("out_type") or "其他",
                   r.get("code") or "", r.get("brand") or "",
                   r.get("category") or "", r.get("model") or "", r.get("shelf") or "",
                   qty, round(cost, 2), round(cost * qty, 2), r.get("remark") or ""])
    _style_sheet(ws, headers, len(rows), widths)
    for idx in (8, 9, 10):
        letter = get_column_letter(idx)
        for i in range(2, len(rows) + 2):
            ws["%s%d" % (letter, i)].alignment = Alignment(horizontal="right")
        if idx >= 9:
            for i in range(2, len(rows) + 2):
                ws["%s%d" % (letter, i)].number_format = "#,##0.00"
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)

    # ---- 汇总表：分类型 + 合计 ----
    s = wb.create_sheet("出库汇总")
    s.cell(row=1, column=1, value="指标").font = _HDR_FONT
    s.cell(row=1, column=1).fill = _HDR_FILL
    s.cell(row=1, column=2, value="数值").font = _HDR_FONT
    s.cell(row=1, column=2).fill = _HDR_FILL
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 20
    s.cell(row=2, column=1, value="统计范围").font = _BODY_FONT
    s.cell(row=2, column=2, value=scope_text or "全部").font = _BODY_FONT

    qty_all = sum(int(r.get("qty") or 0) for r in rows)
    cost_all = sum(float(r.get("cost") or 0) * int(r.get("qty") or 0) for r in rows)
    n_sale = sum(1 for r in rows if (r.get("out_type") or OUT_TYPE_SALE) == OUT_TYPE_SALE)
    n_other = len(rows) - n_sale
    qty_sale = sum(int(r.get("qty") or 0) for r in rows
                   if (r.get("out_type") or OUT_TYPE_SALE) == OUT_TYPE_SALE)
    r0 = 3
    for label, val in (("出库笔数(全部)", len(rows)), ("出库件数(全部)", qty_all),
                       ("出库成本价值(元)", round(cost_all, 2)),
                       ("其中销售出库笔数", n_sale), ("其中销售出库件数", qty_sale),
                       ("其中非销售出库笔数", n_other),
                       ("其中非销售出库件数", qty_all - qty_sale)):
        s.cell(row=r0, column=1, value=label).font = _BODY_FONT
        c = s.cell(row=r0, column=2, value=val)
        c.font = _BODY_FONT
        if isinstance(val, float):
            c.number_format = "#,##0.00"
        r0 += 1

    by_type = {}
    for r in rows:
        k = r.get("out_type") or "其他"
        d = by_type.setdefault(k, {"count": 0, "qty": 0, "cost": 0.0})
        d["count"] += 1
        d["qty"] += int(r.get("qty") or 0)
        d["cost"] += float(r.get("cost") or 0) * int(r.get("qty") or 0)
    r0 += 1
    s.cell(row=r0, column=1, value="分类型统计").font = Font(
        name="微软雅黑", size=10.5, bold=True)
    r0 += 1
    for i, h in enumerate(("类型", "笔数", "件数", "成本价值(元)"), start=1):
        c = s.cell(row=r0, column=i, value=h)
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(horizontal="center")
        c.border = _BORDER
    for col, w in (("A", 26), ("B", 12), ("C", 12), ("D", 16)):
        s.column_dimensions[col].width = w
    for k in sorted(by_type):
        v = by_type[k]
        r0 += 1
        s.cell(row=r0, column=1, value=k).font = _BODY_FONT
        s.cell(row=r0, column=2, value=v["count"]).font = _BODY_FONT
        s.cell(row=r0, column=3, value=v["qty"]).font = _BODY_FONT
        c = s.cell(row=r0, column=4, value=round(v["cost"], 2))
        c.font = _BODY_FONT
        c.number_format = "#,##0.00"
    wb.save(path)
    return path


PURCHASE_HEADERS = [
    ("purchase_date", "采购日期", 12), ("code", "编号", 14), ("brand", "品牌", 16),
    ("category", "品类", 12), ("model", "型号", 19), ("shelf", "货架", 15),
    ("qty", "采购数量", 10), ("unit_cost", "采购单价(元)", 13),
    ("total", "采购金额(元)", 13), ("supplier", "供应商", 14),
    ("remark", "备注", 20),
]


def export_purchases_xlsx(path, rows, scope_text=""):
    """导出采购记录：明细表 + 汇总（合计 + 分供应商）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "采购记录"
    headers = [h[1] for h in PURCHASE_HEADERS]
    widths = {i + 1: h[2] for i, h in enumerate(PURCHASE_HEADERS)}
    ws.append(headers)
    for r in rows:
        ws.append([r.get("purchase_date") or "", r.get("code") or "",
                   r.get("brand") or "", r.get("category") or "",
                   r.get("model") or "", r.get("shelf") or "",
                   int(r.get("qty") or 0), round(float(r.get("unit_cost") or 0), 2),
                   round(float(r.get("total") or 0), 2), r.get("supplier") or "",
                   r.get("remark") or ""])
    _style_sheet(ws, headers, len(rows), widths)
    for idx in (7, 8, 9):
        letter = get_column_letter(idx)
        for i in range(2, len(rows) + 2):
            ws["%s%d" % (letter, i)].alignment = Alignment(horizontal="right")
        if idx >= 8:
            for i in range(2, len(rows) + 2):
                ws["%s%d" % (letter, i)].number_format = "#,##0.00"
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)

    s = wb.create_sheet("采购汇总")
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 20
    s.cell(row=1, column=1, value="指标").font = _HDR_FONT
    s.cell(row=1, column=1).fill = _HDR_FILL
    s.cell(row=1, column=2, value="数值").font = _HDR_FONT
    s.cell(row=1, column=2).fill = _HDR_FILL
    s.cell(row=2, column=1, value="统计范围").font = _BODY_FONT
    s.cell(row=2, column=2, value=scope_text or "全部").font = _BODY_FONT

    qty_all = sum(int(r.get("qty") or 0) for r in rows)
    amt_all = sum(float(r.get("total") or 0) for r in rows)
    r0 = 3
    for label, val in (("采购笔数", len(rows)), ("采购件数", qty_all),
                       ("采购金额(元)", round(amt_all, 2)),
                       ("平均进价(元)", round(amt_all / qty_all, 2) if qty_all else 0.0)):
        s.cell(row=r0, column=1, value=label).font = _BODY_FONT
        c = s.cell(row=r0, column=2, value=val)
        c.font = _BODY_FONT
        if isinstance(val, float):
            c.number_format = "#,##0.00"
        r0 += 1

    sup = {}
    for r in rows:
        k = (r.get("supplier") or "").strip() or "（未填写）"
        slot = sup.setdefault(k, {"count": 0, "qty": 0, "amount": 0.0})
        slot["count"] += 1
        slot["qty"] += int(r.get("qty") or 0)
        slot["amount"] += float(r.get("total") or 0)
    r0 += 1
    s.cell(row=r0, column=1, value="分供应商统计").font = Font(
        name="微软雅黑", size=10.5, bold=True)
    r0 += 1
    for i, h in enumerate(("供应商", "笔数", "件数", "采购金额(元)"), start=1):
        c = s.cell(row=r0, column=i, value=h)
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(horizontal="center")
        c.border = _BORDER
    for col, w in (("A", 26), ("B", 12), ("C", 12), ("D", 16)):
        s.column_dimensions[col].width = w
    for k in sorted(sup, key=lambda x: -sup[x]["amount"]):
        v = sup[k]
        r0 += 1
        s.cell(row=r0, column=1, value=k).font = _BODY_FONT
        s.cell(row=r0, column=2, value=v["count"]).font = _BODY_FONT
        s.cell(row=r0, column=3, value=v["qty"]).font = _BODY_FONT
        c = s.cell(row=r0, column=4, value=round(v["amount"], 2))
        c.font = _BODY_FONT
        c.number_format = "#,##0.00"
    wb.save(path)
    return path


def write_template(path):
    """生成导入模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "库存模板"
    headers = [FIELD_LABEL[k] for k in FIELD_KEYS]
    widths = {i + 1: FIELD_WIDTH[k] / 8.2 + 2 for i, k in enumerate(FIELD_KEYS)}
    ws.append(headers)
    samples = [
        ["SKU-1001", "罗技", "外设", "G502 LIGHTSPEED", "99新", 389.00, 12, "A-01", "是", "罗技 G502 无线鼠标 拆封未使用"],
        ["SKU-1002", "Apple", "手机", "A2634", "95新", 2899.00, 3, "B-03", "是", "iPhone 13 128G 国行 无锁"],
        ["SKU-1003", "戴尔", "显示器", "U2720Q", "9成新", 1580.00, 5, "C-02", "否", "戴尔 U2720Q 边角轻微磕碰"],
        ["SKU-1004", "联想", "笔记本", "X1 Carbon Gen9", "8成新", 3600.00, 2, "D-01", "否", "ThinkPad X1C 电池健康 88%"],
    ]
    for s in samples:
        ws.append(s)
    _style_sheet(ws, headers, len(samples), widths)
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(samples) + 1)
    wb.save(path)
    return path
